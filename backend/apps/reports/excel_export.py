from django.db.models import Sum
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from rest_framework import permissions
from rest_framework.views import APIView
from django.utils import timezone

from apps.inventory.models import DailyBalance, DailyStock, StockAdjustment

from .helpers import resolve_showroom_scope


def _styled_header(ws, headers):
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)


class DailyExcelExportView(APIView):
    """
    GET /api/reports/export/daily/?date=YYYY-MM-DD&showroom=<id>
    Downloads one showroom's daily statement (stock + balance) as .xlsx.
    """

    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        target_date = request.query_params.get("date", timezone.localdate().isoformat())
        showrooms_qs, showroom_id = resolve_showroom_scope(request)
        if not showroom_id:
            return HttpResponse("showroom parameter is required", status=400)

        wb = Workbook()

        ws_stock = wb.active
        ws_stock.title = "Stock"
        _styled_header(ws_stock, ["Product", "Opening", "Received", "Sold", "Return", "Adjustment", "Closing"])
        stocks = DailyStock.objects.filter(showroom_id=showroom_id, date=target_date).select_related("product")
        adjustments = StockAdjustment.objects.filter(
            location="showroom", showroom_id=showroom_id, created_at__date=target_date,
        ).select_related("product")
        adjustment_by_product = {}
        for entry in adjustments:
            row = adjustment_by_product.setdefault(entry.product_id, {"difference": 0, "new_quantity": entry.new_quantity})
            row["difference"] += entry.difference
            row["new_quantity"] = entry.new_quantity
        for s in stocks:
            adjustment = adjustment_by_product.pop(s.product_id, {"difference": 0})
            ws_stock.append([s.product.name, s.opening_qty, s.received_qty, s.sold_qty, s.return_qty,
                             adjustment["difference"], s.closing_qty + adjustment["difference"]])
        for product_id, adjustment in adjustment_by_product.items():
            product = next(entry.product for entry in adjustments if entry.product_id == product_id)
            ws_stock.append([product.name, 0, 0, 0, 0, adjustment["difference"], adjustment["new_quantity"]])

        ws_balance = wb.create_sheet("Balance")
        _styled_header(ws_balance, ["Item", "Amount"])
        balance = DailyBalance.objects.filter(showroom_id=showroom_id, date=target_date).first()
        if balance:
            rows = [
                ("Opening Balance", balance.opening_balance),
                ("Cash Sale", balance.cash_sale),
                ("Card Sale", balance.card_sale),
                ("Total Sale", balance.total_sale),
                ("Expense", balance.expense),
                ("Salary", balance.salary),
                ("Deposit", balance.deposit),
                ("Closing Balance", balance.closing_balance),
            ]
            for row in rows:
                ws_balance.append(row)

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        filename = f"daily-statement-{showroom_id}-{target_date}.xlsx"
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        wb.save(response)
        return response


class MonthlyExcelExportView(APIView):
    """
    GET /api/reports/export/monthly/?year=YYYY&month=MM&showroom=<id optional for admin>
    Downloads the monthly balance + stock summary as .xlsx.
    """

    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        year = int(request.query_params.get("year", timezone.localdate().year))
        month = int(request.query_params.get("month", timezone.localdate().month))
        showrooms_qs, showroom_id = resolve_showroom_scope(request)

        balances = DailyBalance.objects.filter(date__year=year, date__month=month, showroom__in=showrooms_qs)
        if showroom_id:
            balances = balances.filter(showroom_id=showroom_id)

        stocks = DailyStock.objects.filter(date__year=year, date__month=month, showroom__in=showrooms_qs)
        if showroom_id:
            stocks = stocks.filter(showroom_id=showroom_id)

        wb = Workbook()
        ws_balance = wb.active
        ws_balance.title = "Balance Summary"
        _styled_header(ws_balance, ["Showroom", "Total Sale", "Cash Sale", "Card Sale", "Expense", "Salary", "Deposit"])
        balance_summary = balances.values("showroom__name").annotate(
            total_sale=Sum("total_sale"), cash_sale=Sum("cash_sale"), card_sale=Sum("card_sale"),
            expense=Sum("expense"), salary=Sum("salary"), deposit=Sum("deposit"),
        )
        for row in balance_summary:
            ws_balance.append([
                row["showroom__name"], row["total_sale"], row["cash_sale"], row["card_sale"],
                row["expense"], row["salary"], row["deposit"],
            ])

        ws_stock = wb.create_sheet("Stock Summary")
        _styled_header(ws_stock, ["Product", "Received", "Sold", "Return", "Adjustment"])
        stock_summary = stocks.values("product__name").annotate(
            received=Sum("received_qty"), sold=Sum("sold_qty"), returned=Sum("return_qty"),
        )
        adjustment_summary = StockAdjustment.objects.filter(
            location="showroom", showroom__in=showrooms_qs,
            created_at__date__year=year, created_at__date__month=month,
        )
        if showroom_id:
            adjustment_summary = adjustment_summary.filter(showroom_id=showroom_id)
        adjustment_by_name = {
            row["product__name"]: row["adjustment"] or 0
            for row in adjustment_summary.values("product__name").annotate(adjustment=Sum("difference"))
        }
        seen_products = set()
        for row in stock_summary:
            name = row["product__name"]
            seen_products.add(name)
            ws_stock.append([name, row["received"], row["sold"], row["returned"], adjustment_by_name.pop(name, 0)])
        for name, difference in adjustment_by_name.items():
            ws_stock.append([name, 0, 0, 0, difference])

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        filename = f"monthly-report-{year}-{month:02d}.xlsx"
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        wb.save(response)
        return response
